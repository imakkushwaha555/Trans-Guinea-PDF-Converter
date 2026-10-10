import io
import tempfile
import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document

from app import app
from pdf_common import TEMP_PREFIX
from tools.convert import _convert_to_visual_docx
from tools.split import parse_ranges
from pdf_common import ConversionError


@pytest.fixture()
def client():
    app.config.update(TESTING=True)
    with app.test_client() as test_client:
        yield test_client


def test_home_page(client):
    response = client.get("/")
    assert response.status_code == 200
    for path in (b"/convert", b"/merge", b"/split"):
        assert path in response.data
    assert b'rel="canonical"' in response.data
    assert b"application/ld+json" in response.data


@pytest.mark.parametrize("path", ["/convert", "/merge", "/split"])
def test_tool_pages_have_seo_metadata(client, path):
    html = client.get(path).get_data(as_text=True)
    assert "<title>" in html and 'name="description"' in html
    assert 'property="og:title"' in html
    assert f'href="http://localhost{path}"' in html


def test_robots_and_sitemap(client):
    assert "Sitemap:" in client.get("/robots.txt").get_data(as_text=True)
    sitemap = client.get("/sitemap.xml").get_data(as_text=True)
    assert "/merge" in sitemap and "/split" in sitemap


def test_security_headers(client):
    response = client.get("/")
    assert response.headers["X-Content-Type-Options"] == "nosniff"
    assert "frame-ancestors 'none'" in response.headers["Content-Security-Policy"]


def test_health(client):
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json == {"status": "ok"}


def test_requires_file(client):
    response = client.post("/convert", data={"format": "docx"})
    assert response.status_code == 400
    assert "choose a PDF" in response.json["error"]


def test_rejects_non_pdf(client):
    response = client.post(
        "/convert",
        data={"format": "docx", "file": (io.BytesIO(b"not a pdf"), "notes.txt")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert "Only PDF" in response.json["error"]


def _make_pdf(password=None, pages=1):
    pdf = pymupdf.open()
    for index in range(pages):
        pdf.new_page().insert_text((72, 90), f"Page {index + 1} sample text", fontsize=14)
    buffer = io.BytesIO()
    if password:
        pdf.save(buffer, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw=password)
    else:
        pdf.save(buffer)
    pdf.close()
    return buffer.getvalue()


def _post(client, data, name="sample.pdf", fmt="docx", password=None):
    form = {"format": fmt, "file": (io.BytesIO(data), name)}
    if password is not None:
        form["password"] = password
    return client.post("/convert", data=form, content_type="multipart/form-data")


def _job_dirs():
    return list(Path(tempfile.gettempdir()).glob(f"{TEMP_PREFIX}*"))


def test_password_protected_pdf_requests_password(client):
    response = _post(client, _make_pdf("secret"), "locked.pdf")
    assert response.status_code == 401
    assert response.json["password_required"] is True
    assert response.json["incorrect"] is False


def test_wrong_password_is_flagged(client):
    response = _post(client, _make_pdf("secret"), "locked.pdf", password="nope")
    assert response.status_code == 401
    assert response.json["incorrect"] is True


def test_correct_password_returns_file_and_cleans_up(client):
    response = _post(client, _make_pdf("secret"), "locked.pdf", password="secret")
    assert response.status_code == 200
    assert response.headers["Content-Disposition"].startswith("attachment")
    assert Document(io.BytesIO(response.data)).paragraphs is not None
    assert _job_dirs() == []


def test_corrupted_pdf_is_rejected_and_cleaned_up(client):
    response = _post(client, b"%PDF-1.7\n1 0 obj\n<< broken", "broken.pdf")
    assert response.status_code == 400
    assert "damaged" in response.json["error"]
    assert _job_dirs() == []


def test_non_pdf_content_with_pdf_name_is_rejected(client):
    response = _post(client, b"plain text, not a pdf", "fake.pdf")
    assert response.status_code == 400
    assert "not a valid PDF" in response.json["error"]
    assert _job_dirs() == []


def test_download_route_is_gone(client):
    assert client.get("/download/anything.docx").status_code == 404


def test_rejects_unknown_format(client):
    response = client.post(
        "/convert",
        data={"format": "pptx", "file": (io.BytesIO(b"%PDF-1.7"), "sample.pdf")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert "supported output format" in response.json["error"]


def test_visual_docx_contains_page_image(tmp_path):
    pdf_path = tmp_path / "visual-source.pdf"
    docx_path = tmp_path / "visual-output.docx"
    pdf = pymupdf.open()
    page = pdf.new_page(width=612, height=792)
    page.insert_text((72, 90), "Visual fidelity test", fontsize=18)
    page.draw_rect(pymupdf.Rect(72, 130, 250, 250), color=(0.5, 0.1, 0.3), fill=(0.9, 0.5, 0.7))
    pdf.save(pdf_path)
    pdf.close()

    _convert_to_visual_docx(Path(pdf_path), Path(docx_path))

    document = Document(docx_path)
    assert len(document.inline_shapes) == 1


def _page_texts(pdf_bytes):
    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as document:
        return [page.get_text("text").strip() for page in document]


def _merge(client, files, passwords=None):
    data = {"files": [(io.BytesIO(content), name) for name, content in files]}
    if passwords is not None:
        import json

        data["passwords"] = json.dumps(passwords)
    return client.post("/merge", data=data, content_type="multipart/form-data")


def test_merge_keeps_upload_order_and_cleans_up(client):
    first, second = _make_pdf(pages=1), _make_pdf(pages=2)
    response = _merge(client, [("a.pdf", first), ("b.pdf", second)])
    assert response.status_code == 200
    assert response.mimetype == "application/pdf"
    assert len(_page_texts(response.data)) == 3
    assert _job_dirs() == []


def test_merge_requires_two_files(client):
    response = _merge(client, [("a.pdf", _make_pdf())])
    assert response.status_code == 400
    assert "at least 2" in response.json["error"]


def test_merge_rejects_non_pdf(client):
    response = _merge(client, [("a.pdf", _make_pdf()), ("b.txt", b"text")])
    assert response.status_code == 400
    assert "Only PDF" in response.json["error"]


def test_merge_names_damaged_file(client):
    response = _merge(client, [("a.pdf", _make_pdf()), ("broken.pdf", b"%PDF-1.7\n<< broken")])
    assert response.status_code == 400
    assert "broken.pdf" in response.json["error"]
    assert _job_dirs() == []


def test_merge_asks_for_password_per_file(client):
    files = [("a.pdf", _make_pdf()), ("locked.pdf", _make_pdf("secret"))]
    response = _merge(client, files)
    assert response.status_code == 401
    assert response.json["password_required"] is True
    assert response.json["file_index"] == 1
    assert response.json["file_name"] == "locked.pdf"

    wrong = _merge(client, files, {"1": "nope"})
    assert wrong.status_code == 401 and wrong.json["incorrect"] is True

    ok = _merge(client, files, {"1": "secret"})
    assert ok.status_code == 200
    assert len(_page_texts(ok.data)) == 2
    assert _job_dirs() == []


def test_parse_ranges():
    assert parse_ranges("1-3, 5, 8-10", 10) == [(1, 3), (5, 5), (8, 10)]


@pytest.mark.parametrize("text", ["", "abc", "0", "3-1", "11", "1-99", "1,,x"])
def test_parse_ranges_rejects_bad_input(text):
    with pytest.raises(ConversionError):
        parse_ranges(text, 10)


def _split(client, data, **fields):
    form = {"file": (io.BytesIO(data), "doc.pdf"), **fields}
    return client.post("/split", data=form, content_type="multipart/form-data")


def test_split_ranges_returns_zip(client):
    response = _split(client, _make_pdf(pages=5), mode="ranges", ranges="1-2, 4")
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        names = sorted(archive.namelist())
        assert names == ["doc_part1_pages_1-2.pdf", "doc_part2_pages_4.pdf"]
        assert len(_page_texts(archive.read(names[0]))) == 2
        assert len(_page_texts(archive.read(names[1]))) == 1
    assert _job_dirs() == []


def test_split_single_range_returns_pdf(client):
    response = _split(client, _make_pdf(pages=4), mode="ranges", ranges="2-3")
    assert response.status_code == 200
    assert response.mimetype == "application/pdf"
    assert len(_page_texts(response.data)) == 2


def test_split_every_page(client):
    response = _split(client, _make_pdf(pages=3), mode="every_page")
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        assert len(archive.namelist()) == 3


def test_split_rejects_out_of_range(client):
    response = _split(client, _make_pdf(pages=2), mode="ranges", ranges="1-5")
    assert response.status_code == 400
    assert "outside the document" in response.json["error"]
    assert _job_dirs() == []


def test_split_asks_for_password(client):
    locked = _make_pdf("secret", pages=3)
    assert _split(client, locked, mode="ranges", ranges="1").status_code == 401
    ok = _split(client, locked, mode="ranges", ranges="1", password="secret")
    assert ok.status_code == 200
    assert _job_dirs() == []

def test_help_page_and_search(client):
    assert client.get("/help").status_code == 200
    assert b'id="helpSearch"' in client.get("/").data
    titles = [item["title"] for item in client.get("/help/search?q=password").get_json()]
    assert "Password-protected PDFs" in titles
    assert client.get("/help/search?q=").get_json() == []
